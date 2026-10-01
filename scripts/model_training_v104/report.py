from __future__ import annotations
from pathlib import Path
import json, math
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DISPLAY={
 'soil_moisture_forecast':'Model 1 - Soil Moisture Forecasting',
 'nutrient_recommendation':'Model 2 - Nutrient Recommendation',
 'anomaly_detection':'Model 3 - Anomaly Detection',
 'smart_irrigation_scheduler':'Model 4 - Smart Irrigation Scheduling',
 'predictive_maintenance':'Model 5 - Predictive Maintenance',
}

def portable(v):
    s=str(v)
    for marker in ['model-artifacts/','logs/','data/','scripts/','docs/']:
        if marker in s:
            return marker+s.split(marker,1)[1]
    return s

def fmt(v):
    if isinstance(v,float): return f'{v:.4f}'
    return str(v)

def shade(cell,fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h); shade(t.rows[0].cells[i]); t.rows[0].cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=fmt(v)
    return t

def add_picture(doc,path,caption,width=6.15):
    if path and Path(path).exists():
        doc.add_picture(str(path),width=Inches(width)); p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.style='Caption'

def add_model_section(doc,m,idx):
    name=m['model']; title=DISPLAY.get(name,name)
    doc.add_heading(f'{idx}. {title}',1)
    doc.add_heading(f'{idx}.1. Mục tiêu nghiệp vụ và bài toán ML',2)
    doc.add_paragraph(m['objective'])
    table(doc,['Thuộc tính','Giá trị'],[
        ['Task',m['task']],['Target',m['target']],['Số feature',len(m['features'])],['Train rows',m['rows']['train']],['Validation rows',m['rows']['validation']],['Test rows',m['rows']['test']]
    ])
    doc.add_paragraph('Feature đầu vào: '+', '.join(m['features'])+'.')
    doc.add_heading(f'{idx}.2. Lý do lựa chọn thuật toán',2)
    p=doc.add_paragraph(); p.add_run('Lập luận trước khi xem kết quả: ').bold=True; p.add_run(m['algorithm_rationale'])
    doc.add_paragraph('Nguyên tắc lựa chọn: tất cả candidate chỉ được so sánh trên Validation. Test set được giữ lại để đánh giá model đã chọn, không dùng để chọn thuật toán.')
    rows=[]
    for c in m['candidate_validation']:
        metrics=', '.join(f'{k}={fmt(v)}' for k,v in c['validation_metrics'].items())
        rows.append([c['candidate'],metrics,'CHỌN' if c['candidate']==m['selected_algorithm'] else 'Không chọn'])
    table(doc,['Candidate','Validation metrics','Kết quả lựa chọn'],rows)
    doc.add_paragraph(f"Thuật toán được chọn: {m['selected_algorithm']}.")
    doc.add_heading(f'{idx}.3. Kết quả Test khóa',2)
    table(doc,['Metric','Giá trị'],[[k,v] for k,v in m['test_metrics'].items()])
    if m['task']=='classification':
        doc.add_paragraph('Class distribution trên Test: '+', '.join(f'{k}={v}' for k,v in m.get('class_distribution',{}).items()))
        rep=m.get('classification_report',{})
        rows=[]
        for k,v in rep.items():
            if isinstance(v,dict) and all(x in v for x in ['precision','recall','f1-score','support']):
                rows.append([k,v['precision'],v['recall'],v['f1-score'],int(v['support'])])
        if rows: table(doc,['Lớp','Precision','Recall','F1','Support'],rows)
    doc.add_heading(f'{idx}.4. Biểu đồ, ma trận và diễn giải',2)
    ch=m.get('charts',{})
    if m['task']=='regression':
        add_picture(doc,ch.get('actual_pred'),'Actual vs Predicted. Điểm càng sát đường chéo càng tốt; khoảng cách tới đường chéo thể hiện sai số dự đoán.')
        doc.add_paragraph('Giải thích: nếu đám mây điểm lệch có hệ thống lên trên hoặc xuống dưới đường chéo, model có bias. Các điểm xa đường chéo là trường hợp model sai lớn cần xem theo zone/thời điểm.')
        add_picture(doc,ch.get('residuals'),'Residual distribution (Prediction - Actual). Residual lý tưởng tập trung quanh 0 và tương đối cân đối.')
        doc.add_paragraph('Giải thích: residual lệch khỏi 0 cho thấy model có xu hướng dự đoán cao/thấp. Đuôi dài thể hiện tồn tại một số trường hợp sai số lớn dù MAE tổng thể thấp.')
        add_picture(doc,ch.get('sequence'),'Chuỗi Actual vs Predicted trên phần cuối Test set, dùng để kiểm tra model có bám được biến động theo thứ tự thời gian hay không.')
    else:
        add_picture(doc,ch.get('confusion'),'Confusion Matrix. Hàng là nhãn thực tế, cột là nhãn dự đoán; đường chéo là dự đoán đúng, ô ngoài đường chéo là lỗi.')
        doc.add_paragraph('Giải thích ma trận: True Positive/True Negative là dự đoán đúng; False Positive tạo cảnh báo giả; False Negative là model bỏ sót sự cố. Với anomaly/predictive maintenance, False Negative cần được ưu tiên giám sát khi chuyển sang dữ liệu thực.')
        add_picture(doc,ch.get('roc'),'ROC curve cho bài toán nhị phân. Đường cong càng áp sát góc trên-trái càng thể hiện khả năng tách lớp tốt.')
        add_picture(doc,ch.get('pr'),'Precision-Recall curve. Biểu đồ này đặc biệt hữu ích khi lớp sự cố hiếm vì thể hiện trực tiếp đánh đổi giữa bắt đủ sự cố và giảm cảnh báo giả.')
    add_picture(doc,ch.get('feature_importance'),'Feature Importance của model đã chọn. Đây là mức đóng góp tương đối trong mô hình cây, không được diễn giải thành quan hệ nhân quả.')
    if m.get('feature_importance'):
        top=m['feature_importance'][:8]; table(doc,['Feature','Importance'],[[x['feature'],x['importance']] for x in top])
    doc.add_heading(f'{idx}.5. Artifact và khả năng tái lập',2)
    table(doc,['Thuộc tính','Giá trị'],[['Artifact',portable(m['artifact'])],['SHA-256',m['sha256']],['Predictions CSV',portable(m['predictions'])],['Trạng thái','EXPERIMENTAL_SIMULATION']])


def generate(full,out_path:Path):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10); styles['Title'].font.name='Arial'
    title=doc.add_heading('BÁO CÁO HUẤN LUYỆN MÔ HÌNH AI & MÔ PHỎNG END-TO-END - NEXTFARM',0); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f"Phiên bản: {full['version']}\nRun ID: {full['run_id']}\nNgày tạo: {full['created_at']}")
    warn=doc.add_paragraph(); r=warn.add_run('PHẠM VI: DỮ LIỆU MÔ PHỎNG 30 NGÀY - TRAIN THẬT TRÊN DATASET SYNTHETIC. KHÔNG PHẢI BẰNG CHỨNG HIỆU NĂNG NGOÀI THỰC ĐỊA.'); r.bold=True

    doc.add_heading('1. Tóm tắt quy trình thực nghiệm',1)
    doc.add_paragraph('Báo cáo này là tài liệu duy nhất mô tả toàn bộ vòng đời thực nghiệm: sinh/đọc dữ liệu mô phỏng, kiểm tra dữ liệu, tạo feature và target, chia Train/Validation/Test, benchmark thuật toán, huấn luyện thật bằng scikit-learn, đánh giá trên Test khóa, sinh biểu đồ/ma trận, đóng gói model, chạy mô phỏng API shadow, batch inference, A/B routing, monitoring và kiểm tra retraining entrypoint.')
    steps=['Data ingestion từ data/simulation-30day-v103','Data audit + kiểm tra trùng/missing/sự cố','Feature engineering causal + target mô phỏng','Temporal/customer holdout split','Train nhiều candidate cho từng model','Chọn thuật toán bằng Validation','Đánh giá một lần trên Test','Đóng gói .joblib + .pkl kèm preprocessing','Shadow API + batch inference + A/B 10%','Monitoring + retraining readiness']
    for i,s in enumerate(steps,1): doc.add_paragraph(f'{i}. {s}')

    doc.add_heading('2. Dữ liệu và kiểm tra trước train',1); a=full['dataset_audit']
    table(doc,['Chỉ tiêu','Giá trị'],[['Nguồn dữ liệu','SYNTHETIC_SIMULATION_ONLY'],['Rows raw',a['rows_raw']],['Rows sau feature engineering',a['rows_featured']],['Customers',a['customers']],['Devices',a['devices']],['Thời gian bắt đầu',a['start']],['Thời gian kết thúc',a['end']],['Số ngày',a['duration_days']],['Duplicate key rows',a['duplicate_key_rows']]])
    table(doc,['Cảm biến','Tỷ lệ missing'],[[k,v] for k,v in a['missing_fraction'].items()]); table(doc,['Sự cố ground-truth mô phỏng','Positive rows'],[[k,v] for k,v in a['incident_positive_rows'].items()])
    doc.add_paragraph('Nguyên tắc cleaning: không backfill dữ liệu tương lai; trạng thái thiết bị chỉ forward-fill tối đa 2 bước để giữ tính nhân quả. Outlier hợp lệ vật lý không tự động xóa vì có thể chính là tín hiệu sự cố.')

    doc.add_heading('3. Feature engineering, target và chống data leakage',1)
    doc.add_paragraph('Feature lag/rolling chỉ sử dụng quá khứ. Soil Moisture dùng lag 10/30/60/180 phút, rolling mean/std 1 giờ và slope 1 giờ. Các biến giờ được mã hóa sin/cos. Target Soil Moisture là t+6 giờ. Target anomaly/maintenance là có incident trong 60 phút tương lai. Nutrient và irrigation dùng policy synthetic được ghi rõ, không giả định là ground truth thực địa.')
    s=full['split']; table(doc,['Thuộc tính split','Giá trị'],[['Start',s['start']],['Train end',s['train_end']],['Validation end',s['validation_end']],['End',s['end']],['Held-out customers',', '.join(s['heldout_customers'])]])
    doc.add_paragraph('Train và Validation dùng nhóm customer phát triển nhưng tách theo thời gian; Test dùng các customer held-out ở giai đoạn thời gian cuối. Cấu trúc này khó hơn random split và giảm nguy cơ leakage theo thời gian/customer.')

    for i,m in enumerate(full['models'],4): add_model_section(doc,m,i)

    base=4+len(full['models'])
    doc.add_heading(f'{base}. Đóng gói mô hình (Model Packaging)',1)
    doc.add_paragraph('Mỗi model đã train được đóng gói hai định dạng .joblib và .pkl. Artifact chứa sklearn Pipeline bao gồm SimpleImputer và estimator, danh sách feature, target, task, scope, label mapping (nếu có). Vì preprocessing nằm trong Pipeline nên serving không được fit lại scaler/imputer trên dữ liệu inference.')
    table(doc,['Model','Artifact','SHA-256'],[[m['model'],Path(m['artifact']).name,m['sha256'][:24]+'...'] for m in full['models']])

    doc.add_heading(f'{base+1}. Deployment / Serving / System Integration',1)
    d=full['deployment_smoke']; table(doc,['Kiểm thử','Kết quả'],[[k,portable(v)] for k,v in d.items() if not isinstance(v,(dict,list))])
    if d.get('models'):
        table(doc,['Model','HTTP/Inference','Shadow prediction exposed','Latency ms'],[[x['model'],x['status'],x['exposed'],x['latency_ms']] for x in d['models']])
    doc.add_paragraph('Shadow mode vẫn chạy model và ghi log nhưng prediction không được expose như một khuyến nghị hành động. A/B routing dùng hash ổn định của entity_id để chọn khoảng 10% entity, tránh cùng một khu vườn thay đổi nhóm ngẫu nhiên giữa các request.')

    doc.add_heading(f'{base+2}. Batch inference, Shadow, A/B Testing và Monitoring',1)
    b=full['batch_smoke']; table(doc,['Batch check','Giá trị'],[[k,v] for k,v in b.items()])
    mon=full['monitoring']; table(doc,['Monitoring','Giá trị'],[['Inference records',mon.get('records',0)],['Model count',len(mon.get('by_model',{}))]])
    for mn,v in mon.get('by_model',{}).items(): doc.add_paragraph(f"{mn}: requests={v['requests']}, latency mean={v['latency_ms_mean']:.3f} ms, p95={v['latency_ms_p95']:.3f} ms, exposed_fraction={v['exposed_fraction']:.3f}.")
    doc.add_paragraph('A/B 10% ở bản mô phỏng chỉ kiểm tra cơ chế phân nhóm. Khi triển khai thật, KPI A/B phải là metric vận hành/nông học đã định nghĩa trước, không chỉ metric ML offline.')

    doc.add_heading(f'{base+3}. Giám sát và tái huấn luyện',1)
    rt=full['retraining']; table(doc,['Thuộc tính','Giá trị'],[[k,v] for k,v in rt.items()])
    doc.add_paragraph('Retraining production phải quay lại Data Readiness Gate, chỉ lấy nguồn dữ liệu thật đã whitelist. Dataset synthetic tiếp tục được giữ cho regression test và benchmark, không được dùng để tự động thay thế production artifact.')

    doc.add_heading(f'{base+4}. Kết luận thực nghiệm',1)
    doc.add_paragraph('Tất cả 5 model trong báo cáo này đã được thực sự fit bằng scikit-learn trên dataset mô phỏng 30 ngày và đánh giá trên tập Test tách biệt. Các con số không được nhập thủ công: chúng được lấy từ object metric sinh trong cùng run training. Tuy nhiên, vì dữ liệu và một số target do simulator/policy sinh ra, kết quả chỉ chứng minh pipeline ML, packaging và deployment simulation hoạt động end-to-end. Chỉ sau khi train lại bằng telemetry/ground-truth thực địa và qua field pilot mới có thể đánh giá production readiness.')
    doc.add_heading(f'{base+5}. Phụ lục - File train và khả năng tái lập',1)
    for f in full['training_files']: doc.add_paragraph(f)
    doc.add_paragraph('Lệnh tái lập toàn bộ: python scripts/model_training_v104/train_all_and_report.py')
    doc.add_paragraph('Mọi model riêng cũng có thể chạy độc lập với --out <run_dir>.')
    out_path.parent.mkdir(parents=True,exist_ok=True); doc.save(out_path); return out_path
