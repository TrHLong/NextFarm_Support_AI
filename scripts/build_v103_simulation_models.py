from __future__ import annotations
import json, hashlib, shutil, platform
from pathlib import Path
from datetime import datetime, timezone
import numpy as np
import pandas as pd
import joblib, sklearn
from sklearn.pipeline import Pipeline
from sklearn.impute import SimpleImputer
from sklearn.ensemble import ExtraTreesRegressor, ExtraTreesClassifier, RandomForestClassifier
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score, accuracy_score, precision_score, recall_score, f1_score, confusion_matrix, roc_auc_score
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SEED=20260928
H=6 # 60 minutes on 10-min grid
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data/simulation-30day-v103'
OUTROOT=ROOT/'model-artifacts/simulation-v103'
RUN=datetime.now(timezone.utc).strftime('sim30-%Y%m%dT%H%M%SZ')
OUT=OUTROOT/RUN
MODELDIR=OUT/'models'; CHART=OUT/'charts'; PRED=OUT/'predictions'
for p in [MODELDIR,CHART,PRED]: p.mkdir(parents=True,exist_ok=True)

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def savejson(p,x): Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str),encoding='utf-8')

def load():
    frames=[]
    for cdir in sorted((DATA/'customers').glob('synthetic_customer_*')):
        s=pd.read_csv(cdir/'sensor_readings.csv',parse_dates=['observed_at'])
        st=pd.read_csv(cdir/'device_status.csv',parse_dates=['observed_at'])
        tr=pd.read_csv(cdir/'incident_ground_truth.csv',parse_dates=['observed_at'])
        s['observed_at']=pd.to_datetime(s.observed_at,utc=True); st['observed_at']=pd.to_datetime(st.observed_at,utc=True); tr['observed_at']=pd.to_datetime(tr.observed_at,utc=True)
        keys=['customer_id','device_id','observed_at']
        df=s.merge(st,on=keys,how='left',suffixes=('','_status')).merge(tr,on=keys,how='left',suffixes=('','_truth'))
        df=df.sort_values('observed_at').reset_index(drop=True)
        # causal short status fill only; outage itself remains represented by missing/online features
        status_cols=['pump_running','valve_running','fertilizer_running','supply_voltage','rssi','pressure','pump_current','sensor_bus_errors','packet_loss_rate','command_failure_rate','online']
        for c in status_cols:
            if c in df: df[c]=df[c].ffill(limit=2)
        frames.append(df)
    return pd.concat(frames,ignore_index=True)

def features(df):
    df=df.sort_values(['customer_id','observed_at']).copy()
    g=df.groupby('customer_id',sort=False)
    for lag in [1,3,6,18]: df[f'moisture_lag_{lag*10}m']=g.soil_moisture.shift(lag)
    df['moisture_mean_1h']=g.soil_moisture.transform(lambda x:x.rolling(6,min_periods=2).mean())
    df['moisture_std_1h']=g.soil_moisture.transform(lambda x:x.rolling(6,min_periods=2).std())
    df['moisture_slope_1h']=(df.soil_moisture-g.soil_moisture.shift(6))/60
    hr=df.observed_at.dt.hour+df.observed_at.dt.minute/60
    df['hour_sin']=np.sin(2*np.pi*hr/24); df['hour_cos']=np.cos(2*np.pi*hr/24)
    # future targets based only on simulator ground truth/policy, never used as features
    df['target_moisture_6h']=g.soil_moisture.shift(-36)
    for name,cols in {'anomaly_60m':['no_flow','leak'],'maintenance_60m':['sensor_fault','power_loss','mqtt_loss']}.items():
        y=np.zeros(len(df),dtype=int)
        for _,idx in df.groupby('customer_id').groups.items():
            p=df.loc[idx,cols].fillna(0).max(axis=1).to_numpy()
            fut=pd.Series(p[::-1]).rolling(H,min_periods=1).max().to_numpy()[::-1]
            y[np.array(list(idx))]=fut.astype(int)
        df['target_'+name]=y
    # Synthetic agronomy policy: low EC/poor pH -> increase; high EC -> decrease; otherwise keep.
    df['target_nutrient_action']=np.select([(df.ec<1.45)|((df.ph<5.75)&(df.ec<1.8)),(df.ec>2.05)|(df.ph>6.75)], [2,0], default=1) # 0 reduce,1 keep,2 increase
    # Synthetic irrigation policy in minutes, clipped. It is a policy label, not field truth.
    deficit=np.maximum(0,65-df.soil_moisture)
    heat=np.maximum(0,df.temperature-28)
    dry=np.maximum(0,70-df.air_humidity)/20
    df['target_irrigation_minutes']=np.clip(4+0.75*deficit+0.55*heat+1.2*dry,4,35)
    return df

BASE=['soil_moisture','temperature','air_humidity','ec','ph','flow_rate','moisture_lag_10m','moisture_lag_30m','moisture_lag_60m','moisture_lag_180m','moisture_mean_1h','moisture_std_1h','moisture_slope_1h','hour_sin','hour_cos']
DEVICE=['flow_rate','supply_voltage','rssi','pressure','pump_current','sensor_bus_errors','packet_loss_rate','command_failure_rate','pump_running','valve_running','online']
SPECS={
 'soil_moisture_forecast':('regression','target_moisture_6h',BASE,ExtraTreesRegressor(n_estimators=180,min_samples_leaf=3,max_features=.85,n_jobs=-1,random_state=SEED)),
 'nutrient_recommendation':('classification','target_nutrient_action',['ec','ph','soil_moisture','temperature','air_humidity','flow_rate','hour_sin','hour_cos'],RandomForestClassifier(n_estimators=180,min_samples_leaf=3,class_weight='balanced',n_jobs=-1,random_state=SEED)),
 'anomaly_detection':('classification','target_anomaly_60m',DEVICE,ExtraTreesClassifier(n_estimators=200,min_samples_leaf=2,class_weight='balanced',n_jobs=-1,random_state=SEED)),
 'smart_irrigation_scheduler':('regression','target_irrigation_minutes',BASE,ExtraTreesRegressor(n_estimators=180,min_samples_leaf=3,max_features=.85,n_jobs=-1,random_state=SEED)),
 'predictive_maintenance':('classification','target_maintenance_60m',DEVICE,ExtraTreesClassifier(n_estimators=200,min_samples_leaf=2,class_weight='balanced',n_jobs=-1,random_state=SEED)),
}

def split(df):
    customers=sorted(df.customer_id.unique()); held=set(customers[-3:]); dev=~df.customer_id.isin(held); test=df.customer_id.isin(held)
    start,end=df.observed_at.min(),df.observed_at.max(); c1=start+(end-start)*.70; c2=start+(end-start)*.85
    return df[dev&(df.observed_at<c1)],df[dev&(df.observed_at>=c1)&(df.observed_at<c2)],df[test&(df.observed_at>=c2)], {'train_end':c1,'validation_end':c2,'heldout_customers':sorted(held)}

def regression_metrics(y,p): return {'mae':float(mean_absolute_error(y,p)),'rmse':float(mean_squared_error(y,p)**.5),'r2':float(r2_score(y,p))}
def classification_metrics(y,p,proba=None):
    out={'accuracy':float(accuracy_score(y,p)),'precision_macro':float(precision_score(y,p,average='macro',zero_division=0)),'recall_macro':float(recall_score(y,p,average='macro',zero_division=0)),'f1_macro':float(f1_score(y,p,average='macro',zero_division=0))}
    if proba is not None and len(np.unique(y))==2:
        try: out['roc_auc']=float(roc_auc_score(y,proba))
        except: pass
    return out

def charts(name,task,y,p):
    if task=='regression':
        plt.figure(figsize=(6,5)); plt.scatter(y,p,s=8,alpha=.35); lo=min(float(np.min(y)),float(np.min(p))); hi=max(float(np.max(y)),float(np.max(p))); plt.plot([lo,hi],[lo,hi]); plt.xlabel('Actual'); plt.ylabel('Predicted'); plt.title(name+' - Actual vs Predicted'); plt.tight_layout(); plt.savefig(CHART/f'{name}_actual_pred.png',dpi=150); plt.close()
        resid=np.asarray(y)-np.asarray(p); plt.figure(figsize=(6,4)); plt.hist(resid,bins=40); plt.xlabel('Residual'); plt.ylabel('Count'); plt.title(name+' - Residuals'); plt.tight_layout(); plt.savefig(CHART/f'{name}_residuals.png',dpi=150); plt.close()
    else:
        cm=confusion_matrix(y,p); plt.figure(figsize=(5,4)); plt.imshow(cm); plt.title(name+' - Confusion Matrix'); plt.xlabel('Predicted'); plt.ylabel('Actual');
        for (i,j),v in np.ndenumerate(cm): plt.text(j,i,str(v),ha='center',va='center')
        plt.tight_layout(); plt.savefig(CHART/f'{name}_confusion.png',dpi=150); plt.close()

def train():
    raw=load(); df=features(raw)
    tr,va,te,splitmeta=split(df)
    report={'run_id':RUN,'created_at':datetime.now(timezone.utc).isoformat(),'scope':'SYNTHETIC_SIMULATION_ONLY','production_ready':False,'data_origin':'synthetic_device_spec_v11','days':30,'customers':int(df.customer_id.nunique()),'rows':int(len(df)),'split':splitmeta,'environment':{'python':platform.python_version(),'sklearn':sklearn.__version__,'pandas':pd.__version__},'models':{},'warning':'Kết quả chỉ phục vụ đồ án, kiểm thử pipeline và benchmark. Không phải bằng chứng hiệu năng ngoài thực địa.'}
    for name,(task,target,cols,est) in SPECS.items():
        a=tr.dropna(subset=[target]); b=va.dropna(subset=[target]); c=te.dropna(subset=[target]);
        pipe=Pipeline([('imputer',SimpleImputer(strategy='median',add_indicator=True)),('model',est)])
        pipe.fit(a[cols],a[target]); predv=pipe.predict(b[cols]); pred=pipe.predict(c[cols])
        if task=='regression': vm=regression_metrics(b[target],predv); tm=regression_metrics(c[target],pred); proba=None
        else:
            pv=pipe.predict_proba(b[cols]); pt=pipe.predict_proba(c[cols]); vm=classification_metrics(b[target],predv,pv[:,1] if pv.shape[1]==2 else None); tm=classification_metrics(c[target],pred,pt[:,1] if pt.shape[1]==2 else None); proba=pt
        artifact={'pipeline':pipe,'features':cols,'target':target,'task':task,'scope':'SYNTHETIC_SIMULATION_ONLY','status':'EXPERIMENTAL','model_name':name,'run_id':RUN,'label_mapping':({'0':'REDUCE','1':'KEEP','2':'INCREASE'} if name=='nutrient_recommendation' else None)}
        path=MODELDIR/f'{name}.joblib'; joblib.dump(artifact,path,compress=3)
        # portable pickle-format alias for requested packaging; still Python/sklearn dependent.
        shutil.copy2(path,MODELDIR/f'{name}.pkl')
        pd.DataFrame({'actual':c[target].to_numpy(),'prediction':pred}).to_csv(PRED/f'{name}_test_predictions.csv',index=False)
        charts(name,task,c[target].to_numpy(),pred)
        report['models'][name]={'task':task,'target':target,'features':cols,'algorithm':type(est).__name__,'validation':vm,'test':tm,'artifact':str(path.relative_to(ROOT)),'sha256':sha(path),'status':'EXPERIMENTAL_SIMULATION'}
    savejson(OUT/'training_report.json',report)
    pd.DataFrame([{'model':k,**v['test']} for k,v in report['models'].items()]).to_csv(OUT/'model_summary.csv',index=False)
    return report

def make_docx(report):
    d=Document(); styles=d.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10)
    t=d.add_heading('NEXTFARM AI – BÁO CÁO TRAIN MÔ PHỎNG 30 NGÀY',0); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(); r=p.add_run('CẢNH BÁO: '); r.bold=True; p.add_run('Toàn bộ kết quả trong báo cáo này dùng dữ liệu synthetic/simulation. Không được diễn giải là độ chính xác ngoài thực địa hoặc production-ready.')
    d.add_heading('1. Phạm vi dữ liệu',1); d.add_paragraph(f"Nguồn: {report['data_origin']}; thời gian mô phỏng: {report['days']} ngày; số site: {report['customers']}; số dòng sau ghép dữ liệu: {report['rows']:,}.")
    d.add_paragraph('Split: 70% thời gian đầu cho train, 15% tiếp theo cho validation; test dùng 3 customer/site giữ lại và phần thời gian cuối. Cách này giảm leakage so với random row split.')
    d.add_heading('2. Kết quả từng capability',1)
    explanations={
      'soil_moisture_forecast':'Dự báo độ ẩm đất sau 6 giờ. MAE/RMSE càng thấp càng tốt; R² càng gần 1 càng tốt.',
      'nutrient_recommendation':'Phân loại REDUCE/KEEP/INCREASE theo chính sách EC/pH mô phỏng. Đây là policy label của simulator, không phải khuyến cáo nông học đã xác minh.',
      'anomaly_detection':'Phát hiện nguy cơ no-flow/leak trong 60 phút tới từ telemetry thiết bị.',
      'smart_irrigation_scheduler':'Ước lượng số phút tưới theo chính sách mô phỏng dựa trên deficit độ ẩm và điều kiện môi trường.',
      'predictive_maintenance':'Dự báo sensor fault/power loss/MQTT loss trong 60 phút tới.'}
    for name,v in report['models'].items():
        d.add_heading(name,2); d.add_paragraph(explanations[name]); d.add_paragraph('Thuật toán: '+v['algorithm']+'. Trạng thái: '+v['status'])
        tab=d.add_table(rows=1,cols=3); tab.style='Table Grid'; tab.rows[0].cells[0].text='Metric'; tab.rows[0].cells[1].text='Validation'; tab.rows[0].cells[2].text='Test'
        keys=sorted(set(v['validation'])|set(v['test']))
        for k in keys:
            row=tab.add_row().cells; row[0].text=k; row[1].text=f"{v['validation'].get(k,float('nan')):.4f}"; row[2].text=f"{v['test'].get(k,float('nan')):.4f}"
        imgs=list(CHART.glob(name+'_*.png'))
        for img in imgs:
            d.add_picture(str(img),width=Inches(5.9)); cap=d.add_paragraph(img.stem.replace('_',' ')); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if v['task']=='classification': d.add_paragraph('Giải thích ma trận: hàng là nhãn thực tế, cột là dự đoán. Ô ngoài đường chéo là lỗi phân loại; với cảnh báo sự cố cần đặc biệt theo dõi false negative khi chuyển sang dữ liệu thực.')
        else: d.add_paragraph('Biểu đồ Actual vs Predicted càng bám đường chéo càng tốt. Residual nên tập trung quanh 0; lệch có hệ thống cho thấy bias cần xử lý.')
    d.add_heading('3. Đóng gói & triển khai',1); d.add_paragraph('Mỗi model được lưu .joblib và .pkl cùng pipeline imputation, danh sách feature, target, scope và run_id. API FastAPI chỉ phục vụ artifact có scope SYNTHETIC_SIMULATION_ONLY ở chế độ simulation/shadow.')
    d.add_heading('4. Giới hạn và bước chuyển sang dữ liệu thật',1); d.add_paragraph('Các nhãn dinh dưỡng và lịch tưới được tạo bởi policy mô phỏng. Khi có dữ liệu thực, phải chạy lại Data Readiness Gate, thay policy label bằng ground truth/feedback thực tế và đánh giá lại trên test set ngoài thực địa. Không dùng model mô phỏng để tự động điều khiển van/bơm production.')
    path=OUT/'REPORT.docx'; d.save(path); return path

if __name__=='__main__':
    r=train(); p=make_docx(r); print(json.dumps({'run':RUN,'out':str(OUT),'report_docx':str(p),'models':{k:v['test'] for k,v in r['models'].items()}},ensure_ascii=False,indent=2))
