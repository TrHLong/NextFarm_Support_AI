"""Generate a self-contained Vietnamese DOCX training/audit report with charts."""
from pathlib import Path
import json
import numpy as np
import pandas as pd


def _import_deps():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    return plt, Document, Inches, Pt, WD_ALIGN_PARAGRAPH


def _add_kv_table(doc, mapping):
    table=doc.add_table(rows=1,cols=2);table.style="Table Grid"
    table.rows[0].cells[0].text="Chỉ tiêu";table.rows[0].cells[1].text="Giá trị"
    for k,v in mapping.items():
        row=table.add_row().cells;row[0].text=str(k);row[1].text=str(v)
    return table


def generate_training_report(report, out_dir):
    plt,Document,Inches,Pt,ALIGN=_import_deps();out=Path(out_dir);charts=out/"charts";charts.mkdir(exist_ok=True)
    doc=Document();styles=doc.styles;styles["Normal"].font.name="Arial";styles["Normal"].font.size=Pt(10)
    title=doc.add_heading("BÁO CÁO DATA READINESS & TRAINING - NEXTFARM",0);title.alignment=ALIGN.CENTER
    p=doc.add_paragraph();p.alignment=ALIGN.CENTER;p.add_run(f"Run: {report['run_id']}\nPipeline: {report['pipeline_version']}")
    doc.add_heading("1. Mục tiêu và nguyên tắc",1)
    doc.add_paragraph("Pipeline chỉ cho phép dữ liệu thiết bị thật có nguồn gốc truy vết được đi vào production training. Dữ liệu synthetic/simulation/pipeline-test chỉ được dùng để kiểm thử phần mềm và không được tạo model production candidate.")
    doc.add_heading("2. Nguồn dữ liệu",1)
    _add_kv_table(doc,{k:v["path"] for k,v in report.get("sources",{}).items()})
    audit=report["audit"];doc.add_heading("3. Data audit và cleaning",1)
    _add_kv_table(doc,{"Input rows":audit["input_rows"],"Output rows":audit["output_rows"],"Timestamp validity":f"{audit['timestamp_validity']:.2%}","Duplicate fraction":f"{audit['duplicate_fraction']:.2%}","Invalid fraction":f"{audit['invalid_fraction']:.2%}","Real rows":audit["real_rows"],"Non-production rows":audit["non_production_rows"]})
    doc.add_paragraph("Giải thích: dữ liệu sai miền vật lý, timestamp lỗi và quality=bad bị loại; các giá trị bất thường nhưng vẫn hợp lệ vật lý được giữ lại vì có thể chính là tín hiệu sự cố. Không tự động điền trung bình vào dữ liệu cảm biến thiếu.")
    # origins chart
    if audit.get("origins"):
        fig,ax=plt.subplots(figsize=(7,3.5));pd.Series(audit["origins"]).sort_values(ascending=False).plot(kind="bar",ax=ax);ax.set_title("Phân bố data_origin");ax.set_ylabel("Số bản ghi");fig.tight_layout();p=charts/"data_origin.png";fig.savefig(p,dpi=160);plt.close(fig);doc.add_picture(str(p),width=Inches(6.3));doc.add_paragraph("Biểu đồ data_origin dùng để chứng minh nguồn dữ liệu. Production training phải dựa trên real_device/field_device/verified_import; synthetic không được tính vào readiness production.")
    ready=report["readiness"];doc.add_heading("4. Data Readiness Gate",1)
    rows={}
    for k,v in ready["checks"].items(): rows[k]=f"{'PASS' if v['pass'] else 'FAIL'} | value={v.get('value')} | required={v.get('required',v.get('required_max'))}"
    _add_kv_table(doc,rows)
    doc.add_paragraph(f"Kết luận readiness cho Soil Moisture Forecast: {ready['status']}. Nếu FAIL, pipeline dừng và không tạo model production candidate.")
    doc.add_heading("5. Readiness của 5 nhóm AI",1)
    _add_kv_table(doc,{k:v.get("status")+((" - "+v.get("reason")) if v.get("reason") else "") for k,v in ready["model_status"].items()})
    doc.add_heading("6. Horizon coverage",1)
    _add_kv_table(doc,{k+" phút":f"coverage={v['coverage']:.2%}; yêu cầu={v['required']:.2%}; {'PASS' if v['pass'] else 'FAIL'}" for k,v in ready["horizon_coverage"].items()})
    doc.add_heading("7. Kết quả training",1)
    models=report.get("models",[])
    if not models:
        doc.add_paragraph("Không train model trong run này. Đây là hành vi chủ đích khi dữ liệu chưa đạt Data Readiness Gate hoặc pipeline chạy audit-only.")
    for m in models:
        doc.add_heading(f"Soil Moisture Forecast - {m['horizon_minutes']} phút",2)
        _add_kv_table(doc,{"Status":m["status"],"Algorithm":m.get("selected_algorithm"),"Train/Val/Test":m.get("split",{}).get("rows"),"MAE":m.get("test_metrics",{}).get("mae"),"RMSE":m.get("test_metrics",{}).get("rmse"),"R²":m.get("test_metrics",{}).get("r2"),"P90 abs error":m.get("test_metrics",{}).get("p90_absolute_error"),"MAE gain vs persistence":m.get("baseline_mae_gain")})
        pred_path=m.get("prediction_csv")
        if pred_path and Path(pred_path).exists():
            d=pd.read_csv(pred_path);true=f"target_{m['horizon_minutes']}m"
            # Actual vs predicted scatter
            fig,ax=plt.subplots(figsize=(5.4,4));ax.scatter(d[true],d["prediction"],s=8,alpha=.5);lo=min(d[true].min(),d.prediction.min());hi=max(d[true].max(),d.prediction.max());ax.plot([lo,hi],[lo,hi]);ax.set_xlabel("Thực tế");ax.set_ylabel("Dự đoán");ax.set_title("Actual vs Predicted");fig.tight_layout();p=charts/f"actual_pred_{m['horizon_minutes']}.png";fig.savefig(p,dpi=160);plt.close(fig);doc.add_picture(str(p),width=Inches(5.4));doc.add_paragraph("Các điểm càng sát đường chéo càng tốt. Sai lệch có hệ thống khỏi đường chéo cho thấy bias hoặc vùng dữ liệu model chưa học tốt.")
            # residual histogram
            residual=d["prediction"]-d[true];fig,ax=plt.subplots(figsize=(5.4,3.4));ax.hist(residual,bins=35);ax.set_title("Phân bố residual (prediction - actual)");ax.set_xlabel("Residual");fig.tight_layout();p=charts/f"residual_{m['horizon_minutes']}.png";fig.savefig(p,dpi=160);plt.close(fig);doc.add_picture(str(p),width=Inches(5.4));doc.add_paragraph("Residual nên tập trung quanh 0. Lệch hẳn về một phía cho thấy model có xu hướng dự đoán cao hoặc thấp quá mức.")
            # time series sample
            sample=d.tail(min(300,len(d))).copy();fig,ax=plt.subplots(figsize=(7,3.5));ax.plot(sample[true].to_numpy(),label="Actual");ax.plot(sample["prediction"].to_numpy(),label="Predicted");ax.legend();ax.set_title("Chuỗi Actual vs Predicted - phần cuối test set");fig.tight_layout();p=charts/f"timeseries_{m['horizon_minutes']}.png";fig.savefig(p,dpi=160);plt.close(fig);doc.add_picture(str(p),width=Inches(6.4));doc.add_paragraph("Biểu đồ chuỗi giúp phát hiện trường hợp metric tổng đẹp nhưng model phản ứng chậm trước các pha đất khô nhanh hoặc sau tưới.")
    doc.add_heading("8. Chống data leakage",1)
    doc.add_paragraph("Train/Validation/Test được chia theo thời gian. label_end phải nằm hoàn toàn trước biên split kế tiếp, vì vậy target tương lai không được vượt sang Validation/Test. Imputer/scaler nằm trong sklearn Pipeline và chỉ fit bằng Train. Validation dùng để chọn thuật toán; Test chỉ đánh giá model đã chọn và không được dùng để chọn model hay refit.")
    doc.add_heading("9. Hạn chế và cách diễn giải",1)
    doc.add_paragraph("PASS ở Data Readiness chỉ có nghĩa dữ liệu đủ điều kiện kỹ thuật để thử train. CANDIDATE_PASS chỉ có nghĩa model vượt acceptance gate kỹ thuật trên test set đã khóa. Nó không đồng nghĩa đã được kiểm định nông học hoặc an toàn để tự động tưới/châm phân ngoài thực địa. Cần field pilot và phê duyệt nghiệp vụ trước deployment tự động.")
    path=out/"REPORT.docx";doc.save(path);return path
